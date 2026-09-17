"""Build the two MOFI Word deliverables from the demo specification."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'word'
OUT.mkdir(exist_ok=True)
SCRATCH = ROOT / '.document-qa'
SCRATCH.mkdir(exist_ok=True)


def clean(text):
    return re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', text).replace('**', '').replace('`', '')


def heading_text(text):
    text = re.sub(r'^\d+\s+', '', clean(text))
    return re.sub(r'[^\w\s]', ' ', text, flags=re.UNICODE).strip()


def base(title, intro):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.75)
    for style in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        s = d.styles[style]
        s.font.name = 'Calibri'
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.size = Pt(11)
        s.paragraph_format.space_after = Pt(6)
    # The bundled default template may carry a blue title border.
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for border in list(d.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.line_spacing = 1.08
    for style, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
        d.styles[style].font.size = Pt(size)
        d.styles[style].font.bold = True
        d.styles[style].paragraph_format.keep_with_next = True
        d.styles[style].paragraph_format.space_before = Pt(10)
    d.add_paragraph(title, 'Title')
    d.add_paragraph('MOFI  |  Demo đánh giá năng lực  |  Phiên bản 1.0  |  17 tháng 9 năm 2026')
    d.add_paragraph(intro)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('MOFI  •  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    d.core_properties.author = 'MOFI'
    d.core_properties.title = title
    return d


def add_table(d, rows):
    width = 7.0
    n = len(rows[0])
    table = d.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    ratios = {2: [.42, .58], 3: [.20, .39, .41], 4: [.09, .29, .17, .45]}.get(n, [1/n]*n)
    if rows[0][0].strip() == 'Mã' and n == 2:
        ratios = [.12, .88]
    if rows[0][0].strip() == 'Cột':
        ratios = [.25, .23, .52]
    if rows[0][0].strip() == 'Test':
        ratios = [.09, .37, .54]
    for col, ratio in zip(table.columns, ratios):
        col.width = Inches(width*ratio)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        elem = OxmlElement('w:'+side)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            elem.set(qn('w:'+key), val)
        borders.append(elem)
    pr.append(borders)
    for i, values in enumerate(rows):
        row = table.rows[0] if i == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(width*ratios[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.04
            r = p.add_run(clean(value.strip()))
            r.font.size = Pt(9.5)
            if i == 0:
                r.bold = True
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DBEAF5' if i == 0 else ('F5F8FA' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            mar = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:'+side)
                e.set(qn('w:w'), '80')
                e.set(qn('w:type'), 'dxa')
                mar.append(e)
            tcpr.append(mar)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def erd_image():
    im = Image.new('RGB', (1800, 1260), 'white')
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/calibri.ttf', 30)
    small = ImageFont.truetype('C:/Windows/Fonts/calibri.ttf', 25)
    boxes = {
        'users': (680, 35, 1120, 155),
        'portfolios': (45, 290, 485, 410),
        'instruments': (1315, 290, 1755, 410),
        'transactions': (680, 510, 1120, 640),
        'market_prices': (1315, 510, 1755, 640),
        'manual_assets': (45, 780, 485, 900),
        'goals': (680, 780, 1120, 900),
        'watchlist_items': (1315, 780, 1755, 900),
        'alert_rules': (45, 1060, 485, 1180),
        'notifications': (495, 1060, 915, 1180),
        'tasks': (925, 1060, 1305, 1180),
        'learning_progress': (1315, 1060, 1755, 1180),
    }
    def line(points, label, at):
        draw.line(points, fill='#65798A', width=4)
        draw.text(at, label, font=small, fill='black')
    line([(680,95),(265,95),(265,290)], '1 : 1', (285,195))
    line([(265,410),(265,575),(680,575)], '1 : N', (410,537))
    line([(1535,410),(1535,510)], '1 : N', (1550,448))
    line([(1315,350),(1190,350),(1190,575),(1120,575)], '1 : N', (1205,446))
    draw.text((570, 235), 'user_id -> users.id', font=font, fill='#243A4D')
    draw.text((520, 704), 'All lower tables: users 1 : N', font=font, fill='#243A4D')
    draw.text((490, 946), 'Watchlist and alerts also reference instruments', font=small, fill='#243A4D')
    for name, rect in boxes.items():
        draw.rectangle(rect, fill='#F0F6FA', outline='#8298A8', width=2)
        x1,y1,x2,y2 = rect
        draw.text((x1+18,y1+20), name, font=font, fill='black')
        note = 'id PK' if name in ['users','instruments'] else ('instrument_id FK' if name=='market_prices' else 'user_id FK')
        draw.text((x1+18,y1+65), note, font=small, fill='#33485A')
    path = SCRATCH / 'demo-erd.png'
    im.save(path)
    return path


def append_md(d, text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            is_erd = 'mermaid' in line
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            if is_erd:
                p = d.add_paragraph()
                r = p.add_run()
                pic = r.add_picture(str(erd_image()), width=Inches(6.9))
                pic._inline.docPr.set('descr', 'ERD MOFI gồm users, portfolio, transactions và các bảng dữ liệu demo')
                d.add_paragraph('Quan hệ tài chính chính và các bảng dữ liệu theo người dùng. Chi tiết khóa ngoại nằm trong từ điển dữ liệu.')
            else:
                for c in code:
                    d.add_paragraph(c)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = lines[i].strip().strip('|').split('|')
                if not all(re.fullmatch(r'\s*:?-+:?\s*', c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(d, rows)
            continue
        elif line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title = heading_text(line[level:].strip())
            d.add_heading(title, max(1, min(level - 1, 3)))
        elif line.startswith('- '):
            d.add_paragraph(clean(line[2:]), 'List Bullet')
        elif re.match(r'^\d+\. ', line):
            d.add_paragraph(clean(re.sub(r'^\d+\. ', '', line)), 'List Number')
        else:
            d.add_paragraph(clean(line))
        i += 1


def section(text, start, end=None):
    a = text.index(start)
    b = text.index(end, a) if end else len(text)
    return text[a:b]


spec = (ROOT/'docs/demo/SPECIFICATION.md').read_text(encoding='utf-8')
db = (ROOT/'docs/demo/DATABASE.md').read_text(encoding='utf-8')
delivery = (ROOT/'docs/demo/DELIVERY.md').read_text(encoding='utf-8')

overview = base('Giới thiệu dự án MOFI', 'MOFI là website quản lý tài sản và đầu tư cá nhân dùng dữ liệu mô phỏng, được chuẩn bị cho buổi đánh giá năng lực phát triển trong hai ngày. Tài liệu giải thích mục tiêu, công nghệ, các màn hình và mức chức năng dự kiến để người đánh giá biết rõ phần nào có xử lý thật và phần nào là mô phỏng.')
append_md(overview, section(spec, '## 1 ', '## 4 '))
append_md(overview, section(spec, '## 6 '))
append_md(overview, section(delivery, '## 1 ', '## 3 '))
overview.save(OUT/'MOFI_Gioi_thieu_du_an.docx')

technical = base('Đặc tả use case và database MOFI', 'Tài liệu xác định luồng sử dụng, quy tắc tính toán và cấu trúc PostgreSQL cho bản demo hai ngày. Đăng nhập và lưu dữ liệu sẽ hoạt động thật trên dữ liệu giả. Schema này được dùng cho đợt triển khai sắp tới; tại thời điểm lập tài liệu chưa có migration nghiệp vụ hoặc chức năng đã nghiệm thu.')
append_md(technical, section(spec, '## 4 ', '## 6 '))
technical.add_page_break()
append_md(technical, db)
technical.add_page_break()
append_md(technical, section(delivery, '## 3 '))
technical.save(OUT/'MOFI_Use_case_va_Database.docx')
print('Created two Word documents from docs/demo sources.')
